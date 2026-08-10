import io
from docx import Document
from app.parsers.base import ParseResult


def parse_docx(content: bytes) -> ParseResult:
    result = ParseResult()
    try:
        doc = Document(io.BytesIO(content))
        text_parts = []

        for para in doc.paragraphs:
            if para.text:
                text_parts.append(para.text)

        for section in doc.sections:
            for header in [section.header, section.first_page_header, section.even_page_header]:
                if header:
                    for para in header.paragraphs:
                        if para.text:
                            text_parts.append(para.text)
            for footer in [section.footer, section.first_page_footer, section.even_page_footer]:
                if footer:
                    for para in footer.paragraphs:
                        if para.text:
                            text_parts.append(para.text)

        tables = []
        for table in doc.tables:
            cleaned = []
            for row in table.rows:
                cleaned_row = []
                for cell in row.cells:
                    cleaned_row.append(cell.text.strip())
                cleaned.append(cleaned_row)
            if cleaned:
                tables.append(cleaned)

        result.text = "\n".join(text_parts)
        result.tables = tables
        result.metadata["paragraphs"] = len(doc.paragraphs)
        result.metadata["tables"] = len(doc.tables)
    except Exception as e:
        result.errors.append(f"DOCX: {e}")
    return result
