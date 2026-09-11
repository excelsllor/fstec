"""API Gateway (ТЗ 2.6, 3.1). Внешний контур микросервисной архитектуры.

Эндпоинты ТЗ 2.6: POST /upload, GET /documents, GET /documents/{id},
GET /reports/{id}/download, GET /reports/{id}/download_raw, POST /reply/generate.
Плюс обратная совместимость: /api/letters/... и /api/auth/*.
"""
import base64
import logging
import os
import shutil
import threading
import time
from collections import deque
from contextlib import asynccontextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

from fastapi import Depends, FastAPI, File, Form, HTTPException, Request, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, Response
from sqlalchemy.orm import Session

from shared import bus as eventbus
from shared.auth import (
    create_access_token, ensure_bootstrap, get_current_user,
    hash_password, needs_setup, require_admin, verify_password,
)
from shared.config import (
    ALLOWED_EXTENSIONS, BOOTSTRAP_USERNAME, MAX_FILE_SIZE,
    MAX_FILES_PER_BATCH, ORG_NAME, UPLOAD_DIR,
)
from shared.db import SessionLocal, get_db, init_db
from shared.events import DocumentsUploaded
from shared.models import (
    Attachment, AuditLog, BootstrapSecret, Document, Entity, GeneratedResponse,
    IoC, Measure, MeasureCandidate, Report, Summary, Threat, User, Vulnerability,
)

logger = logging.getLogger("fstec.gateway")

# --- Rate limit на /api/auth/login (анти-брутфорс; off для автотестов) ---

_LOGIN_LIMIT = int(os.environ.get("FSTEC_LOGIN_RATE_LIMIT", "10"))
_LOGIN_WINDOW_S = int(os.environ.get("FSTEC_LOGIN_RATE_WINDOW_S", "60"))
_LOGIN_LIMIT_DISABLED = os.environ.get("FSTEC_DISABLE_LOGIN_RATE_LIMIT", "0") == "1"

_login_attempts: dict[str, deque] = {}
_login_lock = threading.Lock()


def _login_allowed(client_ip: str) -> bool:
    if _LOGIN_LIMIT_DISABLED:
        return True
    now = time.monotonic()
    with _login_lock:
        q = _login_attempts.setdefault(client_ip, deque())
        while q and now - q[0] > _LOGIN_WINDOW_S:
            q.popleft()
        if len(q) >= _LOGIN_LIMIT:
            return False
        q.append(now)
        return True

@asynccontextmanager
async def lifespan(app: FastAPI):
    init_db()
    with SessionLocal() as sess:
        pw = ensure_bootstrap(sess)
        if pw:
            logger.warning("Bootstrap admin password: %s", pw)
    yield


app = FastAPI(title="FSTEC API Gateway", version="1.0.0",
              docs_url="/api/docs", openapi_url="/api/openapi.json",
              lifespan=lifespan)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


def _utcnow():
    return datetime.now(timezone.utc)


def _content_disposition(filename: str) -> str:
    """Content-Disposition с ASCII-fallback и RFC 5987 (*=UTF-8'') для кириллицы."""
    from urllib.parse import quote
    try:
        filename.encode("latin-1")
        return f"attachment; filename={filename}"
    except UnicodeEncodeError:
        safe = filename.encode("ascii", "replace").decode("ascii") or "download"
        return f"attachment; filename={safe}; filename*=UTF-8''{quote(filename)}"


def _audit(db: Session, user: User | None, action: str, object_type: str = "",
           object_id: int | None = None, filename: str = "", document_id: int | None = None,
           request: Request | None = None):
    rec = AuditLog(
        user_id=user.id if user else None,
        username=user.username if user else "",
        action=action,
        object_type=object_type,
        object_id=object_id,
        filename=filename,
        document_id=document_id,
        ip_address=(request.client.host if request and request.client else ""),
    )
    db.add(rec)


# ---------------- Auth ----------------

@app.post("/api/auth/login")
def login(body: dict, request: Request, db: Session = Depends(get_db)):
    client_ip = request.client.host if request.client else "unknown"
    if not _login_allowed(client_ip):
        raise HTTPException(status_code=429, detail="Слишком много попыток входа, повторите позже")
    username = (body.get("username") or "").strip()
    password = body.get("password") or ""
    user = db.query(User).filter(User.username == username).first()
    if not user or not verify_password(password, user.password_hash):
        raise HTTPException(status_code=401, detail="Неверный логин или пароль")
    _audit(db, user, "login", "user", user.id, request=request)
    db.commit()
    token = create_access_token({"sub": user.username, "role": user.role})
    return {"access_token": token, "token_type": "bearer", "role": user.role, "username": user.username}


@app.get("/api/auth/status")
def auth_status(db: Session = Depends(get_db)):
    setup = needs_setup(db)
    return {"needs_setup": setup}


@app.get("/api/auth/me")
def auth_me(user: User = Depends(get_current_user)):
    return {"id": user.id, "username": user.username, "role": user.role,
            "full_name": user.full_name, "is_active": user.is_active}


# ---------------- Upload (ТЗ 2.6 POST /upload) ----------------

_MAGIC = {
    ".pdf": (b"%PDF", 4),
    ".docx": (b"PK\x03\x04", 4),
    ".xlsx": (b"PK\x03\x04", 4),
    ".odt": (b"PK\x03\x04", 4),
    ".doc": (b"\xd0\xcf\x11\xe0", 4),
    ".rtf": (b"{\\rtf", 5),
}


def _validate_file(uf: UploadFile):
    ext = Path(uf.filename or "").suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=415, detail=f"Недопустимый формат: {ext or '(нет расширения)'}")
    uf.file.seek(0, os.SEEK_END)
    size = uf.file.tell()
    uf.file.seek(0)
    if size > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail=f"Файл {uf.filename} больше 20 МБ")
    magic, n = _MAGIC.get(ext, (None, 0))
    if magic:
        head = uf.file.read(n)
        uf.file.seek(0)
        if ext == ".rtf":
            head = head.lstrip(b"\x00\x09\x0a\x0d\x20")
        if not head.startswith(magic):
            raise HTTPException(status_code=415, detail=f"Файл {uf.filename} не соответствует формату {ext}")
    uf.file.seek(0)
    return size


@app.post("/upload", status_code=201)
async def upload(
    files: Optional[list[UploadFile]] = File(default=None),
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    if not files:
        raise HTTPException(status_code=400, detail="Нет файлов")
    if len(files) > MAX_FILES_PER_BATCH:
        raise HTTPException(status_code=413, detail=f"Допускается до {MAX_FILES_PER_BATCH} файлов одновременно")

    sizes = []
    for uf in files:
        sizes.append(_validate_file(uf))

    main_name = Path(files[0].filename).name
    doc = Document(
        letter_type="other", status="uploaded", processing_stage="uploaded",
        source_filename=main_name,
        size=sum(sizes), mime=Path(main_name).suffix,
        created_by=user.id,
    )
    db.add(doc)
    db.flush()

    doc_dir = UPLOAD_DIR / f"doc_{doc.id}"
    doc_dir.mkdir(parents=True, exist_ok=True)
    doc.file_path = str(doc_dir)

    att_metas = []
    for idx, uf in enumerate(files):
        name = Path(uf.filename).name
        dest = doc_dir / name
        with open(dest, "wb") as f:
            shutil.copyfileobj(uf.file, f)
        db.add(Attachment(
            document_id=doc.id, filename=name, file_path=str(dest),
            file_type=Path(name).suffix.lower(), parse_status="pending", is_main=(idx == 0),
        ))
        att_metas.append({"filename": name, "path": str(dest), "size": sizes[idx]})

    _audit(db, user, "upload", "document", doc.id, files[0].filename, doc.id)
    db.commit()

    await eventbus.publish("documents.uploaded", str(doc.id), DocumentsUploaded(
        document_id=doc.id,
        user_id=user.id,
        main_file=att_metas[0],
        attachments=att_metas[1:],
    ).model_dump(mode="json"))

    return {"id": doc.id, "status": "uploaded", "filename": files[0].filename}


# ---------------- Documents (ТЗ 2.6) ----------------

@app.get("/documents")
def list_documents(
    skip: int = 0, limit: int = 50,
    letter_type: str | None = None,
    status: str | None = None,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    query = db.query(Document)
    if letter_type:
        query = query.filter(Document.letter_type == letter_type)
    if status:
        query = query.filter(Document.status == status)
    docs = query.order_by(Document.created_at.desc()).offset(skip).limit(limit).all()
    return [{
        "id": d.id, "letter_number": d.letter_number, "letter_date": d.letter_date,
        "letter_type": d.letter_type, "status": d.status, "sla": d.sla, "routing": d.routing,
        "source_filename": d.source_filename, "processing_stage": d.processing_stage,
        "created_at": d.created_at,
    } for d in docs]


def _document_brief(doc: Document):
    return {
        "id": doc.id,
        "letter_number": doc.letter_number,
        "letter_date": doc.letter_date,
        "letter_type": doc.letter_type,
        "status": doc.status,
        "sla": doc.sla,
        "routing": doc.routing,
        "source_filename": doc.source_filename,
        "processing_stage": doc.processing_stage,
        "created_at": doc.created_at,
    }


def _serialize_document(db: Session, doc: Document) -> dict:
    attachments = [{
        "id": a.id, "filename": a.filename, "file_type": a.file_type,
        "parse_status": a.parse_status, "parse_errors": a.parse_errors,
        "parsed_text": a.parsed_text, "is_main": a.is_main,
    } for a in doc.attachments]

    threats = [{
        "id": t.id, "number": t.number, "group_name": t.group_name,
        "threat_type": t.threat_type, "theme": t.theme,
        "archive_name": t.archive_name, "exe_name": t.exe_name,
        "malware_type": t.malware_type, "description": t.description,
        "measures": t.measures,
    } for t in doc.threats]

    iocs = [{
        "id": i.id, "ioc_type": i.ioc_type, "value": i.value,
        "context": i.context, "llm_validated": i.llm_validated, "source": i.source,
    } for i in doc.iocs]

    vulnerabilities = [{
        "id": v.id, "bdu_id": v.bdu_id, "cve_id": v.cve_id,
        "description": v.description, "software": v.software, "severity": v.severity,
        "cvss_score": v.cvss_score, "cpe": v.cpe, "affected_range": v.affected_range,
        "fixed_version": v.fixed_version, "patch_url": v.patch_url,
        "cmdb_match": v.cmdb_match, "current_version": v.current_version,
        "target_version": v.target_version, "source": v.source,
        "recommendation": v.recommendation,
        "is_applicable": v.is_applicable, "applicability_notes": v.applicability_notes,
        "action_type": v.action_type, "action_details": v.action_details,
        "response_point": v.response_point,
    } for v in doc.vulnerabilities]

    entities = [{
        "id": e.id, "entity_type": e.entity_type, "value": e.value, "context": e.context,
    } for e in db.query(Entity).filter(Entity.document_id == doc.id).all()]

    summaries = [{"id": s.id, "summary": s.summary, "confidence": s.confidence}
                 for s in db.query(Summary).filter(Summary.document_id == doc.id).all()]

    resp = (db.query(GeneratedResponse).filter(GeneratedResponse.document_id == doc.id)
            .order_by(GeneratedResponse.id.desc()).first())

    data = _document_brief(doc)
    data.update({
        "subject": doc.subject,
        "original_text": doc.original_text,
        "all_text": doc.all_text,
        "parse_errors": doc.parse_errors,
        "ocr_used": doc.ocr_used,
        "updated_at": doc.updated_at,
        "attachments": attachments,
        "threats": threats,
        "iocs": iocs,
        "vulnerabilities": vulnerabilities,
        "entities": entities,
        "summaries": summaries,
        "ioc_count": len(iocs),
        "vuln_count": len(vulnerabilities),
        "threat_count": len(threats),
        "response": {"text": resp.content or "", "exists": resp is not None, "id": resp.id} if resp
        else {"text": "", "exists": False},
    })
    return data


@app.get("/documents/{doc_id}")
def get_document(doc_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    doc = db.query(Document).filter(Document.id == doc_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Документ не найден")
    return _serialize_document(db, doc)


# ---------------- Reports (ТЗ 2.6) ----------------

@app.get("/reports/{doc_id}/download")
def download_report(doc_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    report = (db.query(Report).filter(Report.document_id == doc_id)
              .order_by(Report.generated_at.desc()).first())
    if not report:
        raise HTTPException(status_code=404, detail="Отчёт ещё не сформирован")
    return FileResponse(report.file_path, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        filename=report.filename)


@app.get("/reports/{doc_id}/download_raw")
def download_raw(doc_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    doc = db.query(Document).filter(Document.id == doc_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Документ не найден")
    main = (db.query(Attachment).filter(Attachment.document_id == doc_id, Attachment.is_main == True).first())
    path = main.file_path if main else (doc.file_path and os.path.join(doc.file_path, doc.source_filename))
    if not path or not Path(path).exists():
        raise HTTPException(status_code=404, detail="Исходный файл не найден")
    return FileResponse(path, filename=doc.source_filename)


@app.post("/reply/generate")
def reply_generate(doc_id: int, body: dict | None = None, db: Session = Depends(get_db),
                   user: User = Depends(get_current_user)):
    """Принудительная генерация проекта ответа (ТЗ 2.6 /reply/generate)."""
    from shared.generator.templated_reply import generate_reply, persist_candidates
    doc = db.query(Document).filter(Document.id == doc_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Документ не найден")
    threat_rows = (db.query(Threat).filter(Threat.document_id == doc_id)
                   .order_by(Threat.number).all())
    blocks = [{
        "id": t.id, "number": t.number, "threat_type": t.threat_type, "theme": t.theme,
        "group_name": t.group_name, "archive_name": t.archive_name, "exe_name": t.exe_name,
        "malware_type": t.malware_type, "description": t.description,
        "explicit_measures": [m.strip() for m in (t.measures or "").split("\n") if m.strip()],
    } for t in threat_rows]
    vulns = [{"cve_id": v.cve_id, "bdu_id": v.bdu_id, "software": v.software,
              "severity": v.severity, "current_version": v.current_version,
              "target_version": v.target_version, "fixed_version": v.fixed_version,
              "cmdb_match": v.cmdb_match, "recommendation": v.recommendation}
             for v in db.query(Vulnerability).filter(Vulnerability.document_id == doc_id).all()]
    iocs = db.query(IoC).filter(IoC.document_id == doc_id).all()
    addr_count = len({i.value.split(":")[0] for i in iocs if i.ioc_type in ("ip", "domain")})

    rep = generate_reply(db=db, doc_id=doc.id, letter_type=doc.letter_type, blocks=blocks,
                         addr_count=addr_count, active_vulns=vulns)
    if rep["candidates"]:
        for i, c in enumerate(rep["candidates"]):
            rep["candidates"][i]["threat_id"] = (blocks[i]["id"] if i < len(blocks) else None)
        persist_candidates(db, doc.id, rep["candidates"])

    resp = GeneratedResponse(document_id=doc_id, content=rep["text"],
                             edited_content=base64.b64encode(rep["docx"]).decode("ascii"),
                             plan_json=rep["plan_json"], created_by=user.id)
    db.add(resp)
    _audit(db, user, "generate_reply", "document", doc_id, doc.source_filename, doc_id)
    db.commit()
    return {"id": resp.id, "document_id": doc_id, "status": "generated", "text": rep["text"],
            "llm_used": rep["llm_used"], "new_measures": len(rep["candidates"])}


# ---------------- Admin: библиотека мер и очередь кандидатов (LLM) ----------------

@app.get("/admin/measures")
def admin_measures(db: Session = Depends(get_db), user: User = Depends(require_admin)):
    rows = db.query(Measure).order_by(Measure.id).all()
    return [{"id": m.id, "text": m.text, "threat_type": m.threat_type, "tags": m.tags,
             "addr_inflection": m.addr_inflection, "source": m.source} for m in rows]


@app.post("/admin/measures")
def admin_measures_add(body: dict, db: Session = Depends(get_db), user: User = Depends(require_admin)):
    text = (body.get("text") or "").strip()
    if not text:
        raise HTTPException(status_code=400, detail="text обязателен")
    dup = db.query(Measure).filter(Measure.text == text).first()
    if dup:
        return {"id": dup.id, "status": "exists"}
    m = Measure(text=text, threat_type=body.get("threat_type", ""),
                tags=body.get("tags", ""), addr_inflection=bool(body.get("addr_inflection", False)),
                source="admin")
    db.add(m)
    db.commit()
    _audit(db, user, "measure_add", "measure", m.id)
    return {"id": m.id, "status": "added"}


@app.delete("/admin/measures/{measure_id}")
def admin_measures_delete(measure_id: int, db: Session = Depends(get_db), user: User = Depends(require_admin)):
    m = db.query(Measure).filter(Measure.id == measure_id).first()
    if not m:
        raise HTTPException(status_code=404, detail="Мера не найдена")
    db.delete(m)
    db.commit()
    _audit(db, user, "measure_delete", "measure", measure_id)
    return {"status": "deleted"}


@app.get("/admin/measures/candidates")
def admin_candidates(status: str | None = "pending", db: Session = Depends(get_db),
                     user: User = Depends(require_admin)):
    q = db.query(MeasureCandidate)
    if status:
        q = q.filter(MeasureCandidate.status == status)
    rows = q.order_by(MeasureCandidate.id.desc()).all()
    return [{"id": c.id, "text": c.text, "note": c.note, "status": c.status,
             "document_id": c.document_id, "threat_id": c.threat_id,
             "created_at": c.created_at.isoformat() if c.created_at else None,
             "reviewed_at": c.reviewed_at.isoformat() if c.reviewed_at else None}
            for c in rows]


@app.post("/admin/measures/candidates/{candidate_id}/accept")
def admin_candidate_accept(candidate_id: int, db: Session = Depends(get_db),
                           user: User = Depends(require_admin)):
    c = db.query(MeasureCandidate).filter(MeasureCandidate.id == candidate_id).first()
    if not c:
        raise HTTPException(status_code=404, detail="Кандидат не найден")
    if c.status == "approved":
        return {"id": c.id, "status": "already_approved"}
    dup = db.query(Measure).filter(Measure.text == c.text).first()
    if not dup:
        dup = Measure(text=c.text, threat_type="", tags="", addr_inflection=False, source="llm")
        db.add(dup)
    c.status = "approved"
    c.reviewed_by = user.id
    c.reviewed_at = _utcnow()
    db.commit()
    _audit(db, user, "candidate_accept", "measure_candidate", candidate_id)
    return {"id": c.id, "status": "approved", "measure_id": dup.id}


@app.post("/admin/measures/candidates/{candidate_id}/reject")
def admin_candidate_reject(candidate_id: int, db: Session = Depends(get_db),
                           user: User = Depends(require_admin)):
    c = db.query(MeasureCandidate).filter(MeasureCandidate.id == candidate_id).first()
    if not c:
        raise HTTPException(status_code=404, detail="Кандидат не найден")
    c.status = "rejected"
    c.reviewed_by = user.id
    c.reviewed_at = _utcnow()
    db.commit()
    _audit(db, user, "candidate_reject", "measure_candidate", candidate_id)
    return {"id": c.id, "status": "rejected"}


# ---------------- Legacy compat: /api/letters ----------------
# (обратная совместимость для существующего фронтенда до его перехода на ТЗ 2.6)

@app.get("/api/letters")
def legacy_list(skip: int = 0, limit: int = 50, letter_type: str | None = None,
                status: str | None = None, db: Session = Depends(get_db),
                user: User = Depends(get_current_user)):
    query = db.query(Document)
    if letter_type:
        query = query.filter(Document.letter_type == letter_type)
    if status:
        query = query.filter(Document.status == status)
    docs = query.order_by(Document.created_at.desc()).offset(skip).limit(limit).all()
    return [{"id": d.id, "letter_number": d.letter_number, "letter_date": d.letter_date,
             "letter_type": d.letter_type, "status": d.status, "created_at": d.created_at,
             "threat_count": db.query(Threat).filter(Threat.document_id == d.id).count(),
             "ioc_count": db.query(IoC).filter(IoC.document_id == d.id).count(),
             "vuln_count": db.query(Vulnerability).filter(Vulnerability.document_id == d.id).count()}
            for d in docs]


@app.get("/api/letters/stats")
def legacy_stats(db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    total = db.query(Document).count()
    return {
        "total_letters": total,
        "hacker_letters": db.query(Document).filter(Document.letter_type == "hacker").count(),
        "vulnerability_letters": db.query(Document).filter(Document.letter_type == "vulnerability").count(),
        "other_letters": db.query(Document).filter(Document.letter_type == "other").count(),
        "total_threats": db.query(Threat).count(),
        "total_iocs": db.query(IoC).count(),
        "total_vulns": db.query(Vulnerability).count(),
    }


@app.get("/api/letters/{letter_id}")
def legacy_detail(letter_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    doc = db.query(Document).filter(Document.id == letter_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Письмо не найдено")
    return _serialize_document(db, doc)


@app.get("/api/letters/{letter_id}/response/text")
def legacy_response_text(letter_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    resp = db.query(GeneratedResponse).filter(GeneratedResponse.document_id == letter_id) \
        .order_by(GeneratedResponse.id.desc()).first()
    return {"text": (resp.content or "") if resp else "", "exists": resp is not None}


@app.put("/api/letters/{letter_id}/response")
def legacy_response_update(letter_id: int, body: dict, db: Session = Depends(get_db),
                           user: User = Depends(get_current_user)):
    text = (body.get("content") or "").strip()
    resp = db.query(GeneratedResponse).filter(GeneratedResponse.document_id == letter_id) \
        .order_by(GeneratedResponse.id.desc()).first()
    if not resp:
        resp = GeneratedResponse(document_id=letter_id, content=text, created_by=user.id)
        db.add(resp)
    else:
        resp.content = text
        resp.updated_at = _utcnow()
    from shared.generator.response_generator import render_reply_docx
    resp.edited_content = base64.b64encode(render_reply_docx(text)).decode("ascii")
    _audit(db, user, "update_response", "document", letter_id)
    db.commit()
    return {"id": resp.id, "text": text}


@app.get("/api/letters/{letter_id}/response/preview")
def legacy_response_preview(letter_id: int, db: Session = Depends(get_db),
                            user: User = Depends(get_current_user)):
    doc = db.query(Document).filter(Document.id == letter_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Письмо не найдено")
    from shared.generator.reply_preview import build_preview
    threats = db.query(Threat).filter(Threat.document_id == doc.id).order_by(Threat.number).all()
    vulns = db.query(Vulnerability).filter(Vulnerability.document_id == doc.id).all()
    iocs = db.query(IoC).filter(IoC.document_id == doc.id).all()
    preview = build_preview(db=db, document=doc, threats=threats,
                            vulnerabilities=vulns, iocs=iocs)
    return preview


@app.put("/api/letters/{letter_id}/response/measures/{threat_id}")
def legacy_measures_update(letter_id: int, threat_id: int, body: dict,
                           db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    measures = body.get("measures") or []
    threat = db.query(Threat).filter(Threat.id == threat_id, Threat.document_id == letter_id).first()
    if not threat:
        raise HTTPException(status_code=404, detail="Угроза не найдена")
    threat.measures = "\n".join(measures)
    db.commit()
    texts = [m.strip() for m in measures if m.strip()]
    return {"measures": measures, "text": "\n".join(f"  {m}" for m in texts)}


def _update_vulnerability(db: Session, user: User, doc_id: int, vuln_id: int,
                          body: dict) -> dict:
    vuln = db.query(Vulnerability).filter(
        Vulnerability.id == vuln_id, Vulnerability.document_id == doc_id).first()
    if not vuln:
        raise HTTPException(status_code=404, detail="Уязвимость не найдена")
    if "is_applicable" in body and body["is_applicable"] is not None:
        vuln.is_applicable = bool(body["is_applicable"])
    for f in ("applicability_notes", "action_type", "action_details", "response_point", "software"):
        if f in body and body[f] is not None:
            setattr(vuln, f, body[f])
    db.commit()
    _audit(db, user, "update_vulnerability", "vulnerability", vuln_id,
           document_id=doc_id)
    return _serialize_document(db, db.query(Document).filter(Document.id == doc_id).first())


@app.put("/api/letters/{letter_id}/vulnerabilities/{vuln_id}")
def legacy_vuln_update(letter_id: int, vuln_id: int, body: dict,
                       db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    return _update_vulnerability(db, user, letter_id, vuln_id, body)


@app.put("/documents/{doc_id}/vulnerabilities/{vuln_id}")
def vuln_update(doc_id: int, vuln_id: int, body: dict,
                db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    return _update_vulnerability(db, user, doc_id, vuln_id, body)


@app.get("/api/letters/{letter_id}/export/{export_type}")
def legacy_export(letter_id: int, export_type: str, db: Session = Depends(get_db),
                  user: User = Depends(get_current_user)):
    doc = db.query(Document).filter(Document.id == letter_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Письмо не найдено")
    iocs = db.query(IoC).filter(IoC.document_id == doc.id).all()
    text = _build_export_text(export_type, iocs)
    if text is None:
        raise HTTPException(status_code=400, detail=f"Неизвестный тип экспорта: {export_type}")
    filename = _EXPORT_FILENAMES.get(export_type, f"{export_type}.txt")
    return Response(text.encode("utf-8"), media_type="text/plain; charset=utf-8",
                    headers={"Content-Disposition": _content_disposition(filename)})


@app.get("/api/letters/{letter_id}/export/{export_type}/docx")
def legacy_export_docx(letter_id: int, export_type: str, db: Session = Depends(get_db),
                       user: User = Depends(get_current_user)):
    doc = db.query(Document).filter(Document.id == letter_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Письмо не найдено")
    iocs = db.query(IoC).filter(IoC.document_id == doc.id).all()
    content = _build_export_text(export_type, iocs, docx=True)
    if content is None:
        raise HTTPException(status_code=400, detail=f"Неизвестный тип экспорта: {export_type}")
    filename = _EXPORT_DOCX_FILENAMES.get(export_type, f"{export_type}.docx")
    return Response(content, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    headers={"Content-Disposition": _content_disposition(filename)})


def _build_export_text(export_type: str, iocs, docx: bool = False):
    """Возвращает текст или bytes для вида экспорта, либо None для неизвестного типа."""
    if docx:
        from shared.generator.ioc_export_docx import export_ips_docx, export_domains_docx, \
            export_emails_docx, export_all_docx
        builders = {
            "emails": export_emails_docx, "ip_addresses": export_ips_docx,
            "domains": export_domains_docx, "ioc_indicators": export_all_docx,
        }
        fn = builders.get(export_type)
        return fn(iocs) if fn else None

    from shared.generator.ioc_export import export_emails, export_ips, export_domains, export_all
    builders = {
        "emails": export_emails, "ip_addresses": export_ips,
        "domains": export_domains, "ioc_indicators": export_all,
    }
    fn = builders.get(export_type)
    return fn(iocs) if fn else None


_EXPORT_FILENAMES = {
    "emails": "emails.txt",
    "ip_addresses": "ip-адреса на блокировку.txt",
    "domains": "Адреса на блокировку.txt",
    "ioc_indicators": "ioc_indicators.txt",
}

_EXPORT_DOCX_FILENAMES = {
    "emails": "emails.docx",
    "ip_addresses": "ip-адреса на блокировку.docx",
    "domains": "Адреса на блокировку.docx",
    "ioc_indicators": "ioc_indicators.docx",
}


@app.get("/api/letters/{letter_id}/download")
def legacy_download(letter_id: int, db: Session = Depends(get_db), user: User = Depends(get_current_user)):
    resp = db.query(GeneratedResponse).filter(GeneratedResponse.document_id == letter_id) \
        .order_by(GeneratedResponse.id.desc()).first()
    if not resp:
        raise HTTPException(status_code=404, detail="Ответ ещё не сгенерирован")
    import io
    from docx import Document as DocxDocument
    if resp.edited_content:
        content = base64.b64decode(resp.edited_content)
    else:
        buf = io.BytesIO()
        d = DocxDocument()
        d.add_paragraph(resp.content or "")
        d.save(buf)
        content = buf.getvalue()
    return Response(content, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    headers={"Content-Disposition": _content_disposition(f"answer_{letter_id}.docx")})


if __name__ == "__main__":
    import uvicorn
    from shared.config import GATEWAY_PORT
    uvicorn.run(app, host="0.0.0.0", port=GATEWAY_PORT)