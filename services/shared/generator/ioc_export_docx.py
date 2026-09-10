"""Экспорт IOC в DOCX-файлы (ТЗ 2.5.1 / справки на блокировку)."""
import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm


def _create_doc():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(14)
    style.paragraph_format.line_spacing = 1.0
    style.paragraph_format.space_after = Pt(0)
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(3)
        section.right_margin = Cm(1.5)
    return doc


def _save(doc):
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def export_ips_docx(iocs) -> bytes:
    ips = sorted({i.value.split(":")[0] for i in iocs if i.ioc_type == "ip"})
    doc = _create_doc()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("Сетевые индикаторы компрометации")
    run.bold = True
    for ip in ips:
        doc.add_paragraph(ip)
    return _save(doc)


def export_domains_docx(iocs) -> bytes:
    domains = sorted({i.value for i in iocs if i.ioc_type == "domain"})
    doc = _create_doc()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("Обеспечить на уровне сетевых средств защиты информации ограничение обращений к адресам:")
    run.bold = True
    for d in domains:
        doc.add_paragraph(d)
    return _save(doc)


def export_emails_docx(iocs) -> bytes:
    emails = sorted({i.value for i in iocs if i.ioc_type == "email"})
    doc = _create_doc()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run("Адреса электронной почты для блокировки:")
    run.bold = True
    for e in emails:
        doc.add_paragraph(e)
    return _save(doc)


def export_all_docx(iocs) -> bytes:
    doc = _create_doc()

    ips = sorted({i.value.split(":")[0] for i in iocs if i.ioc_type == "ip"})
    if ips:
        p = doc.add_paragraph()
        run = p.add_run("Сетевые индикаторы компрометации")
        run.bold = True
        for ip in ips:
            doc.add_paragraph(ip)
        doc.add_paragraph("")

    domains = sorted({i.value for i in iocs if i.ioc_type == "domain"})
    if domains:
        p = doc.add_paragraph()
        run = p.add_run("Обеспечить на уровне сетевых средств защиты информации ограничение обращений к адресам:")
        run.bold = True
        for d in domains:
            doc.add_paragraph(d)
        doc.add_paragraph("")

    hashes = sorted({i.value for i in iocs if i.ioc_type == "hash"})
    if hashes:
        p = doc.add_paragraph()
        run = p.add_run("Файловые индикаторы компрометации (хэши):")
        run.bold = True
        for h in hashes:
            doc.add_paragraph(h)
        doc.add_paragraph("")

    emails = sorted({i.value for i in iocs if i.ioc_type == "email"})
    if emails:
        p = doc.add_paragraph()
        run = p.add_run("Адреса электронной почты для блокировки:")
        run.bold = True
        for e in emails:
            doc.add_paragraph(e)

    return _save(doc)
