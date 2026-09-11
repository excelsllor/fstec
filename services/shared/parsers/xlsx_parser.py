import io
from openpyxl import load_workbook

from shared.parsers._xml_guard import guard_zip_xml
from shared.parsers.base import ParseResult


def parse_xlsx(content: bytes) -> ParseResult:
    result = ParseResult()
    if not guard_zip_xml(content, what="XLSX", errors=result.errors):
        return result
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True)
        text_parts = []
        tables = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                row_data = []
                for cell in row:
                    val = str(cell).strip() if cell is not None else ""
                    row_data.append(val)
                if any(c for c in row_data):
                    tables.append([row_data])
                    text_parts.extend([c for c in row_data if c])
        result.text = "\n".join(text_parts)
        result.tables = tables
        result.metadata["sheets"] = len(wb.worksheets)
    except Exception as e:
        result.errors.append(f"XLSX: {e}")
    return result