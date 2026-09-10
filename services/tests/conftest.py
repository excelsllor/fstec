import os
import tempfile
from pathlib import Path

_tmp = tempfile.mkdtemp(prefix="fstec_tests_")
os.environ["FSTEC_DATA_DIR"] = _tmp
os.environ["DATABASE_URL"] = f"sqlite:///{Path(_tmp) / 'test.db'}"
os.environ["FSTEC_EVENT_BUS"] = "memory"
os.environ["FSTEC_OCR_ENABLED"] = "false"
os.environ["FSTEC_LLM_PROVIDER"] = "heuristic"
os.environ["FSTEC_SECURITY_MODE"] = "mock"

import pytest  # noqa: E402
from docx import Document as DocxDocument  # noqa: E402

from shared.db import SessionLocal, init_db  # noqa: E402
from shared.models import Attachment, Document  # noqa: E402

init_db()

SAMPLE_TEXT = """Управление ФСТЭК России
Исходящий № 456/1 от 15 марта 2024 г.

Уважаемые коллеги!

По данным ФСТЭК России, хакерской группировкой проводятся целевые атаки с использованием вредоносного программного обеспечения.
Осуществляется рассылка вредоносных сообщений с адресов вида evil[.]com и 203.0.113.10:8080, почта bad@evil.com.
Версия IPv6: 2001:db8:0:0:0:0:0:1. Контактный телефон +7 (900) 123-45-67.

Сообщаем Вам о следующих угрозах:

1. Хакерской группировкой Rare Werewolf осуществляется рассылка фишинговых писем с архивом (архив с наименованием «doc.zip»).
   В случае реализации угрозы возможна установка вредоносного программного обеспечения.
   Необходимо обновить антивирусные базы в течение 3 дней.

2. Хакерской группировкой Cloud Werewolf осуществляется эксплуатация уязвимости BDU:2024-12345 (CVE-2021-44228)
   в программном обеспечении Apache Log4j, уровень опасности по CVSS 9.8 — критический.
   Необходимо произвести обновление до версии 2.17.0 до 01.06.2025.
"""


@pytest.fixture(scope="session")
def data_dir() -> Path:
    return Path(_tmp)


def make_docx_bytes(text: str = SAMPLE_TEXT) -> bytes:
    import io
    buf = io.BytesIO()
    d = DocxDocument()
    for line in text.split("\n"):
        d.add_paragraph(line)
    d.save(buf)
    return buf.getvalue()


def create_uploaded_document(tmp_path, filename: str = "letter.docx", content: bytes | None = None) -> int:
    """Создаёт Document + Attachment (файл на диске), статус uploaded. Возвращает id."""
    content = content if content is not None else make_docx_bytes()
    from shared.config import UPLOAD_DIR
    with SessionLocal() as db:
        doc = Document(source_filename=filename, letter_type="other",
                       status="uploaded", processing_stage="uploaded")
        db.add(doc)
        db.flush()
        doc_dir = UPLOAD_DIR / f"doc_{doc.id}"
        doc_dir.mkdir(parents=True, exist_ok=True)
        fpath = doc_dir / filename
        fpath.write_bytes(content)
        db.add(Attachment(document_id=doc.id, filename=filename, file_path=str(fpath),
                          file_type=".docx", parse_status="pending", is_main=True))
        db.commit()
        return doc.id


@pytest.fixture(scope="session")
def client():
    """Один аутентифицированный TestClient на всю сессию; ручки конвейера регистрируются один раз."""
    from fastapi.testclient import TestClient

    from shared.auth import ensure_bootstrap
    from shared.worker import register_handlers
    from ingest_service.worker import handle_uploaded as ingest_handler
    from llm_service.worker import handle_parsed as llm_handler
    from security_service.worker import handle_analyzed as security_handler
    from reporting_service.worker import handle_assessed as reporting_handler

    with SessionLocal() as db:
        password = ensure_bootstrap(db)
    assert password, "bootstrap-пароль должен быть создан"
    from shared.bus import get_event_bus
    register_handlers(get_event_bus(), {
        "documents.uploaded": ingest_handler,
        "document.parsed": llm_handler,
        "document.analyzed": security_handler,
        "security.assessed": reporting_handler,
    })
    from api_gateway.main import app
    with TestClient(app) as c:
        c.token = None
        r = c.post("/api/auth/login", json={"username": "admin", "password": password})
        assert r.status_code == 200, r.text
        c.token = r.json()["access_token"]
        yield c


def headers(client):
    return {"Authorization": f"Bearer {client.token}"}