import io

from docx import Document as DocxDocument

from shared.generator.indicator_card import (
    build_indicator_card, report_filename, root_domain,
)
from shared.generator.response_generator import (
    generate_reply_text, render_reply_docx,
)


def test_root_domain():
    assert root_domain("evil.example.com") == "example.com"
    assert root_domain("sub.very.long.example.co.uk") == "example.co.uk"
    assert root_domain("plain.ru") == "plain.ru"


def test_indicator_card_builds():
    card = build_indicator_card(
        org_name="Организация Заказчика",
        source_filename="letter.docx",
        document_id=7,
        ips=["203.0.113.10:8080", "203.0.113.10:8080"],
        ipv6=["2001:db8::1"],
        domains=["evil.com"],
        emails=["bad@evil.com"],
        vulnerabilities=[{
            "cve_id": "CVE-2021-44228", "bdu_id": "BDU:2024-12345",
            "software": "Apache Log4j", "severity": "critical",
            "cvss_score": 9.8, "current_version": "2.15.0",
            "target_version": "2.17.0", "cmdb_match": True,
        }],
    )
    doc = DocxDocument(io.BytesIO(card))
    assert "CVE-2021-44228" in "".join(c.text for t in doc.tables for r in t.rows for c in r.cells)


def test_report_filename():
    name = report_filename("letter.docx", report_date=None)
    assert name.startswith("Report_letter.docx_")
    assert name.endswith(".docx")


def test_reply_with_vuln():
    text = generate_reply_text(
        letter_number="456/1", letter_date="2024-03-15", letter_type="vulnerability",
        threats=[], vulnerabilities=[{
            "cve_id": "CVE-2021-44228", "software": "Apache Log4j",
            "severity": "critical", "cmdb_match": True,
            "current_version": "2.15.0", "target_version": "2.17.0",
        }], addr_count=2,
    )
    assert "обновление" in text.lower()


def test_reply_with_phishing_threat():
    text = generate_reply_text(
        letter_number="456/1", letter_date="2024-03-15", letter_type="hacker",
        threats=[{"number": 1, "threat_type": "phishing", "theme": "",
                  "description": "рассылка фишинговых писем"}],
        vulnerabilities=[], addr_count=1,
    )
    assert "песочнице" in text.lower() or "sandbox" in text.lower()


def test_reply_no_threat():
    text = generate_reply_text(
        letter_number="", letter_date="", letter_type="other",
        threats=[], vulnerabilities=[], addr_count=0,
    )
    assert "не подвержено риску" in text


def test_reply_compromise_measures():
    text = generate_reply_text(
        letter_number="9/77", letter_date="2026-04-17", letter_type="compromise",
        threats=[], vulnerabilities=[], addr_count=4,
    )
    assert "не подвержено риску" not in text
    assert "контроль журналов" in text
    assert "внеплановое сканирование" in text


def test_reply_multi_threat_numbering():
    threats = [
        {"number": 1, "threat_type": "phishing", "theme": "Акт сверки",
         "description": "фишинговые рассылки"},
        {"number": 2, "threat_type": "malware_attack", "theme": "",
         "description": "вредоносное ПО типа троян"},
        {"number": 3, "threat_type": "phishing", "theme": "Пояснение",
         "description": "фишинговые рассылки"},
    ]
    text = generate_reply_text(
        letter_number="9/99", letter_date="2026-05-20", letter_type="hacker",
        threats=threats, vulnerabilities=[], addr_count=3,
    )
    assert "1. В целях предотвращения" in text
    assert "2. В целях предотвращения" in text
    assert "3. В целях предотвращения" in text
    assert "связанных с Акт сверки" in text or "связанных с акт сверки" in text.lower()


def test_reply_theme_fwd_cleanup():
    text = generate_reply_text(
        letter_number="9/70", letter_date="2026-04-09", letter_type="hacker",
        threats=[{"number": 1, "threat_type": "phishing",
                  "theme": ":Fwd :Re Новый договор с пересмотренной калькуляцией",
                  "description": "фишинговые рассылки"}],
        vulnerabilities=[], addr_count=2,
    )
    assert ":Fwd" not in text
    assert "связанных с Новый договор с пересмотренной калькуляцией" in text


def test_render_reply_docx():
    content = render_reply_docx("Ответ на письмо\n\nТекст ответа.")
    assert DocxDocument(io.BytesIO(content)).paragraphs
    assert content[:2] == b"PK"