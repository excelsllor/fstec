"""Аудит качества генерации ответов против реальных эталонов (DOCX).

Запуск:
    python tools/audit_quality.py

Сравнивает сгенерированные ответы (из PDF-источников) с реальными ответами:
  - смысловое покрытие (каждый реальный абзац покрыт близким сгенерированным);
  - сущности (BDU, CVE, email, кавычки-наименования, IP);
  - орфографию/пунктуацию.

Известные различия, не считающиеся ошибками, перечислены в ALLOWED_MISS_*
(см. также KNOWN_DIFFS.md).
"""
import sys
import io
import os
import re
from difflib import SequenceMatcher
from types import SimpleNamespace

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(REPO, "backend"))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

from docx import Document
from app.parsers.pdf_parser import parse_pdf
from app.extractor.letter_analyzer import analyze_letter
from app.extractor.vuln_extractor import extract_vulns
from app.generator.response_generator import generate_response
from quality_config import ACTIONS as KNOWN_ACTIONS, MANUAL as MANUAL_VULNS

BASE = os.path.dirname(REPO)

TARGETS = [
    "9-70", "9-77", "9-78", "9-81", "9-85", "9-89", "9-93",
    "9-99", "9-104", "9-107", "9-113", "9-118",
]

# Кавычки-наименования из реального ответа, которых НЕТ в генерации,
# но это согласованные косметические различия (см. KNOWN_DIFFS.md).
ALLOWED_MISS_QUOTES = {
    "9-78": {"«Направляем Вам Исх № ... от ...»"},
    "9-81": {"«Исх № .. от ..»"},
    "9-89": {"«ishod_6726_dolzhnost.zip»"},
    "9-99": {"«:FWD :RE Добрый день. Добавил сведения преамбулу и скорректировал п. 4.6. Больше изменений не вносил»"},
}


def extract_docx_paragraphs(data):
    doc = Document(io.BytesIO(data))
    return [p.text.strip() for p in doc.paragraphs if p.text.strip()]


def real_paras(num):
    p = os.path.join(BASE, f"Otvet_na_pismo_{num}.docx")
    if not os.path.exists(p):
        p = os.path.join(BASE, num, f"Ответ  на письмо {num}.docx")
    return extract_docx_paragraphs(open(p, "rb").read())


def gen_paras(num):
    with open(os.path.join(BASE, num, f"{num}.pdf"), "rb") as f:
        content = f.read()
    pr = parse_pdf(content)
    info = analyze_letter(pr.text)
    letter = SimpleNamespace(
        letter_number=info.letter_number, letter_date=info.letter_date,
        letter_type=info.letter_type, original_text=info.raw_text,
        all_text=info.raw_text, id=None,
    )
    threats = []
    for t in info.threats:
        threats.append(SimpleNamespace(
            number=t.number, group_name=t.group_name, threat_type=t.threat_type,
            theme=t.theme, archive_name=t.archive_name, exe_name=t.exe_name,
            malware_type=t.malware_type, description=t.description, measures="",
        ))
    flat = {bdu: v["action"] for bdu, v in KNOWN_ACTIONS.get(num, {}).items()}
    vinfos = extract_vulns(pr.text, pr.tables)
    vulns = []
    for v in vinfos:
        sw = v.software
        if v.bdu_id in KNOWN_ACTIONS.get(num, {}):
            if "software" in KNOWN_ACTIONS[num][v.bdu_id]:
                sw = KNOWN_ACTIONS[num][v.bdu_id]["software"]
        action = flat.get(v.bdu_id, flat.get(v.cve_id, ""))
        vulns.append(SimpleNamespace(
            bdu_id=v.bdu_id, cve_id=v.cve_id, description=v.description,
            software=sw, severity=v.severity, action_type=action,
        ))
    for bdu, sw, a in MANUAL_VULNS.get(num, []):
        vulns.append(SimpleNamespace(
            bdu_id=bdu, cve_id="", description="", software=sw,
            severity="unknown", action_type=a,
        ))
    docx_bytes = generate_response(letter, threats, vulns, db=None, iocs=[])
    return extract_docx_paragraphs(docx_bytes), info, threats, vulns


def word_ratio(a, b):
    aw = re.findall(r"\w+", a.lower())
    bw = re.findall(r"\w+", b.lower())
    if not aw or not bw:
        return 0.0
    inter = sum(min(aw.count(w), bw.count(w)) for w in set(aw))
    return 2.0 * inter / (len(aw) + len(bw))


# ---------- entities ----------
BDU_RE = re.compile(r"BDU:\d{4}-\d+", re.IGNORECASE)
CVE_RE = re.compile(r"CVE-\d{4}-\d+", re.IGNORECASE)
EMAIL_RE = re.compile(r"[\w.+-]+@[\w.-]+\.[a-zа-я]{2,}", re.IGNORECASE)
QUOTE_RE = re.compile(r"«[^»]+»")
IP_RE = re.compile(r"\b\d{1,3}(?:\.\d{1,3}){3}\b")


def entities(paras):
    text = " ".join(paras)
    return {
        "bdu": {m.upper() for m in BDU_RE.findall(text)},
        "cve": {m.upper() for m in CVE_RE.findall(text)},
        "email": {m.lower() for m in EMAIL_RE.findall(text)},
        "quote": {m for m in QUOTE_RE.findall(text)},
        "ip": set(IP_RE.findall(text)),
    }


# ---------- orthography ----------
def ortho_issues(paras):
    issues = []
    for i, t in enumerate(paras):
        if "  " in t:
            issues.append((i, "double-space", t[:100]))
        for m in re.finditer(r"\s([.,;!?])", t):
            # «№ . от .» / «№ .. от ..» — заглушки реального ответа (см. KNOWN_DIFFS.md)
            if m.group(1) == "." and re.search(r"№\s+\.{1,3}\s+от\s+\.{1,3}", t):
                continue
            issues.append((i, f"space-before-{m.group(1)}", t[:100]))
        for m in re.finditer(r"([.,;:!?])\1", t):
            issues.append((i, "double-punct", t[:100]))
        if t.startswith(" ") or t.endswith(" "):
            issues.append((i, "leading/trailing-space", t[:100]))
    return issues


def semantic_unmatched(gen, real):
    """Реальные абзацы, не покрытые ни одним сгенерированным (точно или близко)."""
    missing = []
    used = [False] * len(gen)
    for ri, rp in enumerate(real):
        best, best_r = 0.0, 0.0
        bi = None
        for gi, gp in enumerate(gen):
            if used[gi]:
                continue
            r = SequenceMatcher(None, rp, gp).ratio()
            wr = word_ratio(rp, gp)
            score = max(r, wr)
            if score > best_r:
                best_r, bi, best = score, gi, r
        if bi is not None and (best >= 0.72 or best_r >= 0.78):
            used[bi] = True
        else:
            missing.append((ri, rp, best_r))
    return missing


report = []
grand = {"real": 0, "covered": 0, "uncovered": 0}
for num in TARGETS:
    gen, info, threats, vulns = gen_paras(num)
    real = real_paras(num)
    ge = entities(gen)
    re_ = entities(real)

    report.append("=" * 72)
    report.append(f"ПИСЬМО {num}: реальных абзацев={len(real)}, сгенерировано={len(gen)}, letter_type={info.letter_type}")
    report.append(f"  угроз={len(threats)}, уязвимостей в БД-выгрузке={len(vulns)}")

    missing = semantic_unmatched(gen, real)
    grand["real"] += len(real)
    grand["uncovered"] += len(missing)
    grand["covered"] += len(real) - len(missing)
    if missing:
        report.append(f"  СМЫСЛОВО НЕ ПОКРЫТЫ ({len(missing)}):")
        for ri, rp, sc in missing:
            report.append(f"    real[{ri}] (score {sc:.2f}): {rp[:180]}")
    else:
        report.append("  СМЫСЛОВОЕ ПОКРЫТИЕ: 100%")

    report.append("  СУЩНОСТИ (недостающие в генерации):")
    any_entity_gap = False
    allowed_quotes = ALLOWED_MISS_QUOTES.get(num, set())
    for key in ("bdu", "cve", "email", "quote", "ip"):
        miss = re_[key] - ge[key]
        if key == "quote":
            miss = miss - allowed_quotes
        if miss:
            any_entity_gap = True
            for v in sorted(miss):
                report.append(f"    MISS {key}: {v}")
    if not any_entity_gap:
        report.append("    все сущности оригинала присутствуют")
    extra = set()
    for key in ("bdu", "cve", "email", "ip"):
        extra |= ge[key] - re_[key]
    if extra:
        report.append(f"  ЛИШНИЕ сущности в генерации: {sorted(extra)}")

    issues = ortho_issues(gen)
    report.append(f"  ОРФОГРАФИЯ/ПУНКТУАЦИЯ: ошибок={len(issues)}")
    for i, kind, snippet in issues[:15]:
        report.append(f"    gen[{i}] {kind}: {snippet!r}")

report.append("=" * 72)
report.append(f"ИТОГО: реальных абзацев={grand['real']}, смыслово покрыто={grand['covered']}/{grand['real']} "
              f"({100.0 * grand['covered'] / grand['real']:.1f}%)")

out = "\n".join(report)
with open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "quality_audit.txt"), "w", encoding="utf-8") as f:
    f.write(out)
print(out)
