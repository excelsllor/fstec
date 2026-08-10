"""Генерация всех эталонных ответов (12 писем) в папку generated_otvety.

Запуск:
    python tools/gen_all.py

Пути вычисляются относительно расположения этого файла:
    tools/  ->  fstec-service/  (репозиторий)
    base    ->  <родитель fstec-service>  (папки 9-{num} и generated_otvety)
"""
import sys
import io
import os
from types import SimpleNamespace

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(REPO, "backend"))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

from docx import Document
from app.parsers.pdf_parser import parse_pdf
from app.extractor.letter_analyzer import analyze_letter
from app.extractor.vuln_extractor import extract_vulns
from app.generator.response_generator import generate_response
from quality_config import ACTIONS, MANUAL

BASE = os.path.dirname(REPO)
out_dir = os.path.join(BASE, "generated_otvety")
os.makedirs(out_dir, exist_ok=True)

LETTERS = [
    "9-70", "9-77", "9-78", "9-81", "9-85", "9-89", "9-93",
    "9-99", "9-104", "9-107", "9-113", "9-118",
]


def extract_docx_paragraphs(data):
    doc = Document(io.BytesIO(data))
    return [p.text.strip() for p in doc.paragraphs if p.text.strip()]


report = []
ok, fail = 0, 0
for num in LETTERS:
    pdf = os.path.join(BASE, num, num + ".pdf")
    if not os.path.exists(pdf):
        report.append(f"{num}: НЕТ PDF")
        fail += 1
        continue
    try:
        with open(pdf, "rb") as f:
            content = f.read()
        pr = parse_pdf(content)
        info = analyze_letter(pr.text)
        letter = SimpleNamespace(
            letter_number=info.letter_number, letter_date=info.letter_date,
            letter_type=info.letter_type, original_text=info.raw_text,
            all_text=info.raw_text, id=None,
        )
        threats = []
        for t in info.threats:
            threats.append(SimpleNamespace(
                number=t.number, group_name=t.group_name, threat_type=t.threat_type,
                theme=t.theme, archive_name=t.archive_name, exe_name=t.exe_name,
                malware_type=t.malware_type, description=t.description, measures="",
            ))
        vinfos = extract_vulns(pr.text, pr.tables)
        vulns = []
        for v in vinfos:
            a = ACTIONS.get(num, {}).get(v.bdu_id, {}).get("action", "")
            sw = v.software
            if v.bdu_id in ACTIONS.get(num, {}):
                if "software" in ACTIONS[num][v.bdu_id]:
                    sw = ACTIONS[num][v.bdu_id]["software"]
            vulns.append(SimpleNamespace(
                bdu_id=v.bdu_id, cve_id=v.cve_id, description=v.description,
                software=sw, severity=v.severity, action_type=a,
            ))
        for bdu, sw, a in MANUAL.get(num, []):
            vulns.append(SimpleNamespace(
                bdu_id=bdu, cve_id="", description="", software=sw,
                severity="unknown", action_type=a,
            ))
        docx_bytes = generate_response(letter, threats, vulns, db=None, iocs=[])
        paras = extract_docx_paragraphs(docx_bytes)
        fname = os.path.join(out_dir, f"Ответ на письмо {num}.docx")
        with open(fname, "wb") as f:
            f.write(docx_bytes)
        bdus = sorted({v.bdu_id for v in vulns if v.bdu_id})
        report.append(f"{num}: OK абзацев={len(paras)} угроз={len(threats)} BDU={bdus} -> {os.path.basename(fname)}")
        ok += 1
    except Exception as e:
        report.append(f"{num}: ОШИБКА {type(e).__name__}: {e}")
        fail += 1

report.append(f"ИТОГО: ok={ok}, fail={fail}, папка={out_dir}")
out = "\n".join(report)
with open(os.path.join(out_dir, "_генерация_отчет.txt"), "w", encoding="utf-8") as f:
    f.write(out)
print(out)
